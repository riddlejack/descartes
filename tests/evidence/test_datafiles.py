"""Data files: CSV and TSV, JSON, XML by suffix, XLSX and XLS (invented content)."""

from __future__ import annotations

import datetime as dt
import io
import json
from collections.abc import Callable
from pathlib import Path

import pytest
from conftest import FakeTransport, Reply

from stemma.evidence import api, limits
from stemma.evidence.datafiles import json_text
from stemma.evidence.errors import CaptureStatus
from stemma.evidence.extract import XLSX_TYPE, extract, sniff_content_type
from stemma.evidence.service import EvidenceService

Factory = Callable[..., EvidenceService]


# CSV and TSV ---------------------------------------------------------------------------

CSV = ('region,year,"value, units",note\r\n'
       'North,2023,"1,250","rose ""sharply""\nafter review"\r\n'
       'South,2023,980.5,\r\n')


def test_csv_kept_verbatim_with_dialect(make_service: Factory, run_dir: Path) -> None:
    t = FakeTransport({"https://stats.example.org/t.csv": Reply(
        200, [("Content-Type", "text/csv; charset=utf-8")], CSV.encode())})
    r = make_service(t).capture("https://stats.example.org/t.csv", "regional table")
    assert r["status"] == "ok" and r["content_type"] == "text/csv"
    text = api.read_text(run_dir, r["receipt_id"])
    assert text == CSV.replace("\r\n", "\n")  # only line endings normalised
    dialect = r["format_details"]["csv_dialect"]
    assert dialect["delimiter"] == "," and dialect["quotechar"] == '"'
    assert dialect["rows"] == 3 and dialect["columns_max"] == 4
    assert api.check_quote(run_dir, r["receipt_id"], 'North,2023,"1,250"')["found"]
    assert not api.check_quote(run_dir, r["receipt_id"], '250"')["found"]  # inside 1,250


def test_semicolon_and_tab_dialects() -> None:
    semi = extract("land;jahr;wert\nNord;2023;1.250,5\nSüd;2023;980\n".encode(), "text/csv")
    assert semi.details["csv_dialect"]["delimiter"] == ";"
    tsv = extract(b"region\tvalue\nNorth\t10\nSouth\t12\n", "text/tab-separated-values")
    assert tsv.details["csv_dialect"]["delimiter"] == "\t"
    assert tsv.text == "region\tvalue\nNorth\t10\nSouth\t12\n"


def test_csv_by_file_name_and_undetected_dialect(make_service: Factory, run_dir: Path) -> None:
    (run_dir / "task" / "pack.csv").write_text("a,b\n1,2\n3,4\n")
    r = make_service().capture_file("task/pack.csv", "user's data")
    assert r["content_type"] == "text/csv" and r["format_details"]["csv_dialect"]["delimiter"] == ","
    one = extract(b"just one column of words\nand another line here\n", "text/csv")
    assert one.status is CaptureStatus.OK and one.text
    assert one.details["csv_dialect"]["delimiter"] is None
    assert any("not_detected" in w for w in one.warnings)


# JSON ----------------------------------------------------------------------------------

RAW_JSON = ('{"series":"GDP","values":[1.10,1e5,12345678901234567890,-0.0,2.50E-3],'
            '"unit":"\\u20ac bn","note":null,"ok":true,"dup":1,"dup":2,"nested":{"a":[]},'
            '"empty":{}}')


def test_json_pretty_printed_keeping_every_literal() -> None:
    text, details, warnings = json_text(RAW_JSON)
    assert warnings == [] and details["pretty_printed"]
    for literal in ("1.10", "1e5", "12345678901234567890", "-0.0", "2.50E-3"):
        assert literal in text  # numbers exactly as written
    assert text.count('"dup"') == 2  # duplicate keys kept
    assert text.index('"series"') < text.index('"values"') < text.index('"unit"')  # source order
    assert '"€ bn"' in text and "null" in text and "true" in text
    assert json_text(RAW_JSON)[0] == text  # deterministic
    assert text.startswith('{\n  "series": "GDP",\n  "values": [\n    1.10,')


def test_json_through_capture(make_service: Factory, run_dir: Path) -> None:
    t = FakeTransport({"https://api.example.org/d": Reply(
        200, [("Content-Type", "application/json")], RAW_JSON.encode())})
    r = make_service(t).capture("https://api.example.org/d", "api")
    assert r["status"] == "ok" and r["extractor"]["name"] == "json-pretty"
    assert r["format_details"]["pretty_printed"] is True
    assert api.check_quote(run_dir, r["receipt_id"], '"unit": "€ bn"')["found"]


@pytest.mark.parametrize("raw, why", [
    ('{"a": [1, 2,', "not valid JSON"),
    ("[" * 300 + "]" * 300, "nesting deeper"),
])
def test_json_not_pretty_printed_kept_verbatim(raw: str, why: str) -> None:
    text, details, warnings = json_text(raw)
    assert text == raw and details["pretty_printed"] is False
    assert any(why in w for w in warnings)


def test_json_lone_surrogate_stays_escaped() -> None:
    text, _, warnings = json_text('{"x": "a\\ud800b"}')
    assert warnings == [] and "\\ud800" in text
    text.encode("utf-8")


def test_vendor_json_and_xml_types() -> None:
    geo = extract(b'{"type":"Feature","properties":{"pop":1250}}', "application/geo+json")
    assert geo.extractor["name"] == "json-pretty" and '"pop": 1250' in (geo.text or "")
    sdmx = extract(b"<?xml version='1.0'?><message><obs value='1250.5'>North</obs>"
                   b"<obs value='980'>South</obs></message>",
                   "application/vnd.sdmx.genericdata+xml")
    assert sdmx.extractor["name"] == "markdownify"
    assert "North" in (sdmx.text or "") and "South" in (sdmx.text or "")


# XLSX and XLS --------------------------------------------------------------------------

def _workbook() -> bytes:
    import openpyxl  # type: ignore[import-untyped]

    book = openpyxl.Workbook()
    first = book.active
    first.title = "Regional totals"
    first.append(["Region", "2023", "2024"])
    first.append(["North", 10, 11.5])
    first.append(["South", 1250, 0.1234])
    first["C3"].number_format = "0.0%"
    first["A5"] = "Note:\tsee\nannex"
    first["B6"] = dt.date(2024, 3, 31)
    first["C7"] = "=B2+B3"
    first.merge_cells("A9:C9")
    first["A9"] = "Merged heading"
    first.row_dimensions[8].hidden = True
    first["A8"] = "Hidden row 777"
    second = book.create_sheet("Notes")
    second["B2"] = "Source: invented survey"
    secret = book.create_sheet("Workings")
    secret["A1"] = "Hidden sheet figure 999"
    secret.sheet_state = "hidden"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_xlsx_every_sheet_with_cell_references(make_service: Factory, run_dir: Path) -> None:
    body = _workbook()
    t = FakeTransport({"https://stats.example.org/t.xlsx": Reply(
        200, [("Content-Type", "application/vnd.ms-excel")], body)})  # mislabelled, as servers do
    r = make_service(t).capture("https://stats.example.org/t.xlsx", "workbook")
    assert r["status"] == "ok" and r["content_type"] == XLSX_TYPE
    assert r["extractor"]["name"] == "stemma-xlsx-stream" and r["format_details"]["sheets"] == 3
    text = api.read_text(run_dir, r["receipt_id"])
    assert "## Sheet 1: Regional totals" in text and "## Sheet 3: Workings" in text
    assert "A1: Region\tB1: 2023\tC1: 2024\n" in text
    assert "A2: North\tB2: 10\tC2: 11.5\n" in text
    assert "B3: 1250\tC3: 0.1234 [format 0.0%]" in text
    assert "A5: Note:\\tsee\\nannex" in text  # tab and line break written, not dropped
    assert "B6: 45382 [format yyyy-mm-dd; date 2024-03-31]" in text  # serial kept, date beside
    assert "C7: [no cached value] [formula =B2+B3]" in text
    assert [o["anchor"] for o in r["outline"]] == ["sheet-1", "sheet-2", "sheet-3"]
    section = EvidenceService(run_dir, cache_dir=None).read_section(r["text_id"], "sheet-2")
    assert "B2: Source: invented survey" in section["text"] and "Workings" not in section["text"]
    kinds = {(x["kind"], text[x["start"]:x["end"]].split("\n")[0][:40]) for x in r["flagged_ranges"]}
    assert ("hidden", "A8: Hidden row 777") in kinds
    assert ("table_merged_cells", "A9: Merged heading") in kinds
    assert any(k == "hidden" and "999" in s for k, s in kinds)
    assert any("no cached value" in w for w in r["warnings"])
    checked = api.check_quote(run_dir, r["receipt_id"], "B3: 1250")
    assert checked["found"] and not checked["warnings"]
    assert not api.check_quote(run_dir, r["receipt_id"], "250\tC3")["found"]
    hidden = api.check_quote(run_dir, r["receipt_id"], "Hidden sheet figure 999")
    assert hidden["found"] and hidden["warnings"]


def test_xlsx_limits_fail_typed(monkeypatch: pytest.MonkeyPatch) -> None:
    body = _workbook()
    monkeypatch.setattr(limits, "SHEET_MAX_UNZIPPED_BYTES", 1000)
    bomb = extract(body, XLSX_TYPE)
    assert bomb.status is CaptureStatus.EXTRACTION_FAILED and "unzipped size" in (bomb.detail or "")
    monkeypatch.setattr(limits, "SHEET_MAX_UNZIPPED_BYTES", 512 * 1024 * 1024)
    monkeypatch.setattr(limits, "SHEET_MAX_CELLS", 3)
    many = extract(body, XLSX_TYPE)
    assert many.status is CaptureStatus.EXTRACTION_FAILED and "cells" in (many.detail or "")


def test_broken_spreadsheets_fail_typed() -> None:
    zip_not_xlsx = extract(b"PK\x03\x04" + b"\x00" * 100, XLSX_TYPE)
    assert zip_not_xlsx.status is CaptureStatus.EXTRACTION_FAILED
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600
    assert sniff_content_type("application/octet-stream", ole, "book.xls") == "application/vnd.ms-excel"
    broken_xls = extract(ole, "application/vnd.ms-excel")
    assert broken_xls.status is CaptureStatus.EXTRACTION_FAILED
    assert broken_xls.extractor["name"] == "xlrd"


def test_xlsx_sniffed_without_a_type(run_dir: Path) -> None:
    body = _workbook()
    assert sniff_content_type(None, body) == XLSX_TYPE
    assert sniff_content_type("application/zip", body) == XLSX_TYPE
    other_zip = io.BytesIO()
    import zipfile

    with zipfile.ZipFile(other_zip, "w") as archive:
        archive.writestr("readme.txt", "not a workbook")
    assert sniff_content_type("application/zip", other_zip.getvalue()) == "application/zip"
    assert json.dumps({"ok": True})
