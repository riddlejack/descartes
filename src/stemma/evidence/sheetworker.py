"""Spreadsheet text in its own process: XLSX read by streaming its XML, XLS by xlrd.

Run as `python -I -m stemma.evidence.sheetworker <xlsx|xls> <max_chars>
<max_cells> <max_unzipped_bytes>` with the file on standard input; writes one JSON
object to standard output, under the parent's recorded time and output limits.

Every sheet, in workbook order, becomes a section:

    ## Sheet 2: Regional totals

    A1: Region	B1: 2023	C1: 2024
    A2: North	B2: 10	C2: 11.5 [format 0.0]

Each line is one row; each non-empty cell is `<reference>: <value>`, cells
separated by a tab, so a quotation can name the cell it rests on.

**Values exactly as the file stores them.** XLSX cells are read from the sheet
XML as text, never through a float: `9007199254740993` and `1.10` stay as written.
Anything rendered is written beside the value, never instead of it: a number
format other than General (`[format 0.0%]`), a date for a date-formatted serial
(`[format yyyy-mm-dd; date 2024-03-31]`, in the workbook's 1900 or 1904 calendar,
with serial 60 named as the 29 February 1900 that exists only in Excel's
calendar), and every formula (`[formula =B2+B3]`, `[array formula =... over
C1:C3]`, `[formula =..., shared over D1:D3]`, `[shared formula, from the group of
D1]`). A formula with no cached value is kept as `[no cached value]` with its
formula. No cell is dropped. In text a backslash, tab or line break is written
`\\\\`, `\\t` or `\\n`, so rows stay one line each. XLS (xlrd 2.0.2) stores numbers
as binary doubles; they are written in their shortest exact form, with the cell's
number format and date beside them in the same way.

**Limits are checked while streaming, before anything is built.** The unzipped
size is checked from the archive's directory; cells and the total area of merged
ranges are counted as the XML streams past, so a workbook declaring huge ranges
fails typed without allocating them.

Flagged ranges: a hidden sheet, hidden rows and every column of a hidden column
range are `hidden`; the top-left cell of a merged range is `table_merged_cells`.
"""

from __future__ import annotations

import datetime as dt
import io
import json
import re
import sys
import zipfile
from collections.abc import Iterator
from decimal import ROUND_FLOOR, Decimal, InvalidOperation
from typing import IO, Any
from xml.etree.ElementTree import iterparse


class TooMuch(Exception):
    pass


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("\t", "\\t").replace("\r\n", "\\n") \
        .replace("\n", "\\n").replace("\r", "\\n")


class _Writer:
    def __init__(self, max_chars: int, max_cells: int) -> None:
        self.parts: list[str] = []
        self.length = 0
        self.cells = 0
        self.max_chars = max_chars
        self.max_cells = max_cells
        self.flagged: list[dict[str, Any]] = []
        self.outline: list[dict[str, Any]] = []
        self.warnings: list[str] = []

    def write(self, text: str) -> tuple[int, int]:
        start = self.length
        self.parts.append(text)
        self.length += len(text)
        if self.length > self.max_chars:
            raise TooMuch(f"text passed the limit of {self.max_chars} characters")
        return start, self.length

    def cell(self) -> None:
        self.cells += 1
        if self.cells > self.max_cells:
            raise TooMuch(f"workbook passed the limit of {self.max_cells} cells")

    def flag(self, start: int, end: int, kind: str) -> None:
        if end > start:
            self.flagged.append({"start": start, "end": end, "kind": kind})

    def text(self) -> str:
        return "".join(self.parts)


def _sheet_heading(w: _Writer, number: int, name: str) -> None:
    if number > 1:
        w.write("\n")
    start = w.length
    w.write(f"## Sheet {number}: {_escape(name)}\n\n")
    w.outline.append({"anchor": f"sheet-{number}", "kind": "sheet", "level": 2,
                      "title": f"Sheet {number}: {name}"[:200], "start": start, "end": None})


def _close_sheet(w: _Writer) -> None:
    w.outline[-1]["end"] = w.length


# Dates -----------------------------------------------------------------------------------

def excel_date(serial: Decimal, date1904: bool) -> str:
    """The date and time a serial means in the workbook's calendar, as ISO text."""
    days = int(serial.to_integral_value(rounding=ROUND_FLOOR))
    seconds = int(((serial - days) * 86400).to_integral_value())
    if date1904:
        base = dt.datetime(1904, 1, 1)
    elif days == 60:
        return "1900-02-29, which exists only in Excel's 1900 calendar"
    elif days < 60:
        base = dt.datetime(1899, 12, 31)
    else:
        base = dt.datetime(1899, 12, 30)
    try:
        moment = base + dt.timedelta(days=days, seconds=seconds)
    except OverflowError:
        return "out of the calendar's range"
    return moment.date().isoformat() if seconds == 0 else moment.isoformat()


def _annotate(raw: str, fmt: str | None, date1904: bool) -> str:
    from openpyxl.styles.numbers import is_date_format  # type: ignore[import-untyped]

    if not fmt or fmt == "General":
        return raw
    if is_date_format(fmt):
        try:
            return f"{raw} [format {_escape(fmt)}; date {excel_date(Decimal(raw), date1904)}]"
        except (InvalidOperation, ValueError, OverflowError):
            return f"{raw} [format {_escape(fmt)}]"
    return f"{raw} [format {_escape(fmt)}]"


# XLSX: streaming reader ------------------------------------------------------------------

def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _attr(elem: Any, name: str) -> str | None:
    for key, value in elem.attrib.items():
        if _local(key) == name:
            return str(value)
    return None


_REF = re.compile(r"^([A-Z]{1,3})([0-9]{1,7})$")


def _col_number(letters: str) -> int:
    number = 0
    for char in letters:
        number = number * 26 + (ord(char) - 64)
    return number


def _col_letters(number: int) -> str:
    letters = ""
    while number:
        number, rest = divmod(number - 1, 26)
        letters = chr(65 + rest) + letters
    return letters


def _area(ref: str) -> int:
    parts = ref.upper().split(":")
    if len(parts) != 2:
        return 1
    a, b = _REF.match(parts[0]), _REF.match(parts[1])
    if a is None or b is None:
        return 1
    cols = abs(_col_number(b.group(1)) - _col_number(a.group(1))) + 1
    rows = abs(int(b.group(2)) - int(a.group(2))) + 1
    return cols * rows


def _events(stream: IO[bytes]) -> Iterator[tuple[str, Any]]:
    yield from iterparse(stream, events=("start", "end"))


def _text_of(elem: Any) -> str:
    """Text of <t> elements below `elem`, leaving out phonetic guides (<rPh>)."""
    parts: list[str] = []

    def walk(node: Any) -> None:
        name = _local(node.tag)
        if name == "rPh":
            return
        if name == "t" and node.text:
            parts.append(node.text)
        for child in node:
            walk(child)

    walk(elem)
    return "".join(parts)


class _Book:
    def __init__(self, archive: zipfile.ZipFile) -> None:
        names = set(archive.namelist())
        self.shared: list[str] = []
        if "xl/sharedStrings.xml" in names:
            with archive.open("xl/sharedStrings.xml") as stream:
                for event, elem in _events(stream):
                    if event == "end" and _local(elem.tag) == "si":
                        self.shared.append(_text_of(elem))
                        elem.clear()
        self.formats: dict[int, str] = {}
        self.xf_formats: list[int] = []
        if "xl/styles.xml" in names:
            with archive.open("xl/styles.xml") as stream:
                in_cell_xfs = False
                for event, elem in _events(stream):
                    name = _local(elem.tag)
                    if name == "cellXfs":
                        in_cell_xfs = event == "start"
                    elif event == "end" and name == "numFmt":
                        fmt_id = _attr(elem, "numFmtId") or "0"
                        if fmt_id.isdigit():
                            self.formats[int(fmt_id)] = _attr(elem, "formatCode") or ""
                    elif event == "end" and name == "xf" and in_cell_xfs:
                        fmt_id = _attr(elem, "numFmtId") or "0"
                        self.xf_formats.append(int(fmt_id) if fmt_id.isdigit() else 0)
        rels: dict[str, str] = {}
        if "xl/_rels/workbook.xml.rels" in names:
            with archive.open("xl/_rels/workbook.xml.rels") as stream:
                for event, elem in _events(stream):
                    if event == "end" and _local(elem.tag) == "Relationship":
                        target = _attr(elem, "Target") or ""
                        target = target.lstrip("/") if target.startswith("/") else "xl/" + target
                        rels[_attr(elem, "Id") or ""] = target
        self.sheets: list[tuple[str, str, str]] = []  # (name, state, part path)
        self.date1904 = False
        with archive.open("xl/workbook.xml") as stream:
            for event, elem in _events(stream):
                name = _local(elem.tag)
                if event == "end" and name == "workbookPr":
                    self.date1904 = (_attr(elem, "date1904") or "").lower() in ("1", "true")
                elif event == "end" and name == "sheet":
                    path = rels.get(_attr(elem, "id") or "", "")
                    self.sheets.append((_attr(elem, "name") or "",
                                        _attr(elem, "state") or "visible", path))

    def number_format(self, style: str | None) -> str | None:
        from openpyxl.styles.numbers import BUILTIN_FORMATS

        if style is None or not style.isdigit():
            return None
        index = int(style)
        if index >= len(self.xf_formats):
            return None
        fmt_id = self.xf_formats[index]
        found: str | None = self.formats.get(fmt_id) or BUILTIN_FORMATS.get(fmt_id)
        return found


def _cell_text(book: _Book, cell: Any, ref: str,
               masters: dict[str, str]) -> tuple[str | None, bool]:
    """The cell's entry (without its reference), and whether its formula has no cached
    value. None for a cell holding nothing."""
    kind = _attr(cell, "t") or "n"
    value: str | None = None
    formula = ""
    inline: str | None = None
    for child in cell:
        name = _local(child.tag)
        if name == "v":
            value = child.text or ""
        elif name == "f":
            ftype = _attr(child, "t") or "normal"
            ftext = _escape(child.text or "")
            fref = _attr(child, "ref")
            si = _attr(child, "si")
            if ftype == "array":
                formula = f" [array formula ={ftext} over {fref}]"
            elif ftype == "shared" and ftext:
                if si is not None:
                    masters[si] = ref
                formula = f" [formula ={ftext}, shared over {fref}]"
            elif ftype == "shared":
                master = masters.get(si or "", f"si={si}")
                formula = f" [shared formula, from the group of {master}]"
            elif ftype == "dataTable":
                formula = f" [data table formula over {fref}]"
            else:
                formula = f" [formula ={ftext}]"
        elif name == "is":
            inline = _text_of(child)
    if value == "" and kind not in ("str", "inlineStr"):
        value = None  # an empty <v/> holds no cached value (writers emit one for formulas)
    if value is None and inline is None:
        if formula:
            return "[no cached value]" + formula, True
        return None, False
    if kind == "s":
        try:
            shown = _escape(book.shared[int(value or "")])
        except (ValueError, IndexError):
            shown = f"[shared string {value} missing]"
    elif kind == "inlineStr":
        shown = _escape(inline or "")
    elif kind == "b":
        shown = {"1": "TRUE", "0": "FALSE"}.get(value or "", _escape(value or ""))
    elif kind in ("str", "e", "d"):
        shown = _escape(value or "")
    else:
        shown = _annotate(_escape(value or ""), book.number_format(_attr(cell, "s")),
                          book.date1904)
    return shown + formula, False


def _xlsx(body: bytes, w: _Writer, max_unzipped: int) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(body))
    except zipfile.BadZipFile as exc:
        raise ValueError(f"not a valid XLSX file: {exc}") from exc
    unzipped = sum(info.file_size for info in archive.infolist())
    if unzipped > max_unzipped:
        raise TooMuch(f"unzipped size {unzipped} passes the limit of {max_unzipped} bytes")
    names = set(archive.namelist())
    if "xl/workbook.xml" not in names:
        raise ValueError("not an XLSX workbook (no xl/workbook.xml)")
    book = _Book(archive)
    merged_area = 0
    uncached = 0
    for number, (name, state, path) in enumerate(book.sheets, start=1):
        _sheet_heading(w, number, name)
        body_start = w.length
        if path not in names:
            w.warnings.append(f"sheet {number} ({name}) has no worksheet part {path!r}")
            _close_sheet(w)
            continue
        hidden_cols: set[int] = set()
        entries: dict[str, tuple[int, int]] = {}
        masters: dict[str, str] = {}
        row_number = 0
        with archive.open(path) as stream:
            for event, elem in _events(stream):
                if event != "end":
                    continue
                tag = _local(elem.tag)
                if tag == "col" and (_attr(elem, "hidden") or "") in ("1", "true"):
                    low = int(_attr(elem, "min") or 0)
                    high = int(_attr(elem, "max") or 0)
                    if high - low > w.max_cells:
                        raise TooMuch("a hidden column range passes the cell limit")
                    hidden_cols.update(range(low, high + 1))
                elif tag == "mergeCell":
                    ref = (_attr(elem, "ref") or "").upper()
                    merged_area += _area(ref)
                    if merged_area > w.max_cells:
                        raise TooMuch(f"merged ranges pass the limit of {w.max_cells} cells")
                    top = ref.split(":")[0]
                    if top in entries:
                        w.flag(*entries[top], "table_merged_cells")
                elif tag == "row":
                    row_attr = _attr(elem, "r")
                    row_number = int(row_attr) if row_attr and row_attr.isdigit() else row_number + 1
                    hidden_row = (_attr(elem, "hidden") or "") in ("1", "true")
                    column = 0
                    pieces = 0
                    for cell in elem:
                        if _local(cell.tag) != "c":
                            continue
                        w.cell()
                        ref_attr = (_attr(cell, "r") or "").upper()
                        match = _REF.match(ref_attr)
                        column = _col_number(match.group(1)) if match else column + 1
                        ref = ref_attr if match else f"{_col_letters(column)}{row_number}"
                        shown, no_cache = _cell_text(book, cell, ref, masters)
                        if shown is None:
                            continue
                        uncached += no_cache
                        if pieces:
                            w.write("\t")
                        start, end = w.write(f"{ref}: {shown}")
                        entries[ref] = (start, end)
                        pieces += 1
                        if hidden_row or column in hidden_cols:
                            w.flag(start, end, "hidden")
                    if pieces:
                        w.write("\n")
                    elem.clear()
        if state != "visible":
            w.flag(body_start, w.length, "hidden")
            w.warnings.append(f"sheet {number} ({name}) is {state}")
        _close_sheet(w)
    if uncached:
        w.warnings.append(f"{uncached} formula cells have no cached value")


# XLS -----------------------------------------------------------------------------------------

def _xls_number(value: float) -> str:
    if value.is_integer() and abs(value) < 2**53:
        return str(int(value))
    return repr(value)


def _xls_value(value: float, fmt: str | None, is_date: bool, datemode: int) -> str:
    """An XLS number (a binary double) in its shortest exact form, with its format,
    and its date when the format is a date format, beside it."""
    raw = _xls_number(value)
    if not fmt or fmt == "General":
        return raw
    if is_date:
        return (f"{raw} [format {_escape(fmt)}; date "
                f"{excel_date(Decimal(repr(value)), datemode == 1)}]")
    return f"{raw} [format {_escape(fmt)}]"


def _xls(body: bytes, w: _Writer) -> None:
    import xlrd  # type: ignore[import-untyped]

    try:
        book = xlrd.open_workbook(file_contents=body, formatting_info=True)
    except (NotImplementedError, xlrd.XLRDError):
        book = xlrd.open_workbook(file_contents=body)
    xf_list = getattr(book, "xf_list", [])
    format_map = getattr(book, "format_map", {})

    def format_of(cell: Any) -> str | None:
        index = getattr(cell, "xf_index", None)
        if index is None or index >= len(xf_list):
            return None
        fmt = format_map.get(xf_list[index].format_key)
        return str(fmt.format_str) if fmt is not None else None

    merged_area = 0
    for number, sheet in enumerate(book.sheets(), start=1):
        _sheet_heading(w, number, sheet.name)
        body_start = w.length
        hidden_rows = {r for r, info in getattr(sheet, "rowinfo_map", {}).items() if info.hidden}
        hidden_cols = {c for c, info in getattr(sheet, "colinfo_map", {}).items() if info.hidden}
        merged = getattr(sheet, "merged_cells", [])
        for r0, r1, c0, c1 in merged:
            merged_area += (r1 - r0) * (c1 - c0)
            if merged_area > w.max_cells:
                raise TooMuch(f"merged ranges pass the limit of {w.max_cells} cells")
        merged_tops = {(r0, c0) for r0, _r1, c0, _c1 in merged}
        for row in range(sheet.nrows):
            pieces = 0
            for col in range(sheet.ncols):
                cell = sheet.cell(row, col)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                w.cell()
                if cell.ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_DATE):
                    value = _xls_value(float(cell.value), format_of(cell),
                                       cell.ctype == xlrd.XL_CELL_DATE, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = "TRUE" if cell.value else "FALSE"
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    value = xlrd.error_text_from_code.get(cell.value, f"#ERROR{cell.value}")
                else:
                    value = _escape(str(cell.value))
                if pieces:
                    w.write("\t")
                start, end = w.write(f"{xlrd.colname(col)}{row + 1}: {value}")
                pieces += 1
                if row in hidden_rows or col in hidden_cols:
                    w.flag(start, end, "hidden")
                if (row, col) in merged_tops:
                    w.flag(start, end, "table_merged_cells")
            if pieces:
                w.write("\n")
        if getattr(sheet, "visibility", 0):
            w.flag(body_start, w.length, "hidden")
            w.warnings.append(f"sheet {number} ({sheet.name}) is hidden")
        _close_sheet(w)
    w.warnings.append("xls formulas are not read: xlrd gives cached values only")


def run(kind: str, body: bytes, max_chars: int, max_cells: int,
        max_unzipped: int) -> dict[str, Any]:
    w = _Writer(max_chars, max_cells)
    if kind == "xlsx":
        _xlsx(body, w, max_unzipped)
    elif kind == "xls":
        _xls(body, w)
    else:
        raise ValueError(f"unknown kind {kind}")
    return {"text": w.text(), "outline": w.outline, "flagged": w.flagged,
            "warnings": w.warnings, "sheets": len(w.outline), "cells": w.cells}


def main() -> None:
    kind, max_chars, max_cells, max_unzipped = sys.argv[1], *map(int, sys.argv[2:5])
    body = sys.stdin.buffer.read()
    try:
        result: dict[str, Any] = {"ok": run(kind, body, max_chars, max_cells, max_unzipped)}
    except TooMuch as exc:
        result = {"error": "too_much", "detail": str(exc)}
    except Exception as exc:  # noqa: BLE001 - reported to the parent as a typed failure
        result = {"error": "spreadsheet", "detail": f"{type(exc).__name__}: {exc}"}
    sys.stdout.write(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()
